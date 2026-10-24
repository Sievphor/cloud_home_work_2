import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header / University
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni1 = p_uni.add_run("សាកលវិទ្យាល័យ សៅស៍អុីសថ៍អេយេសៀ\n")
    r_uni1.font.name = 'Khmer OS Muol Light'
    r_uni1.font.size = Pt(13)
    r_uni1.bold = True
    r_uni1.font.color.rgb = RGBColor(16, 44, 87)

    r_uni2 = p_uni.add_run("UNIVERSITY OF SOUTH-EAST ASIA\n")
    r_uni2.font.name = 'Arial'
    r_uni2.font.size = Pt(12)
    r_uni2.bold = True
    r_uni2.font.color.rgb = RGBColor(16, 44, 87)

    r_fac = p_uni.add_run("Faculty of Science and Technology\n\n")
    r_fac.font.name = 'Arial'
    r_fac.font.size = Pt(11)
    r_fac.font.color.rgb = RGBColor(71, 85, 105)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hw = p_title.add_run("Homework 2\n")
    r_hw.font.name = 'Arial'
    r_hw.font.size = Pt(18)
    r_hw.bold = True
    r_hw.font.color.rgb = RGBColor(37, 99, 235)

    r_t = p_title.add_run("CI/CD WITH JENKINS, GITHUB, DOCKER, AMAZON ECR AND DOCKER SWARM ON AWS EC2\n")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(13)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(30, 41, 59)

    # Info Box Table
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Course:", "Advanced Cloud Technology"),
        ("Lecturer:", "Mr. Krol Noeurn"),
        ("Student Name:", "TONN Sievphor"),
        ("GitHub Repository:", "https://github.com/Sievphor/cloud_home_work_2.git"),
        ("Amazon ECR URI:", "908157891706.dkr.ecr.us-east-1.amazonaws.com/sievphor-homework2"),
        ("Academic Year:", "2026 - 2027")
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = 'Arial'
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].runs[0].font.name = 'Arial'
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")

    doc.add_page_break()

    def add_section_h1(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body_p(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_figure_box(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[{title} - សូម Paste រូបភាព Screenshot នៅទីនេះ]")
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(9.5)
        r.italic = True
        r.font.color.rgb = RGBColor(100, 116, 139)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)

    # 1. Introduction
    add_section_h1("១. សេចក្តីផ្តើម (Introduction)")
    add_body_p(
        "គម្រោង Homework 2 នេះមានគោលបំណងបង្កើតប្រព័ន្ធ Continuous Integration និង Continuous Deployment (CI/CD) "
        "ពេញលេញមួយនៅលើបរិស្ថាន Cloud (AWS) ដោយប្រើប្រាស់បច្ចេកវិទ្យា GitHub, Jenkins, Docker, Amazon ECR និង Docker Swarm នៅលើ Amazon EC2។ "
        "នៅពេលអ្នកអភិវឌ្ឍន៍ (Developer) Push Code ថ្មីទៅកាន់ GitHub Repository នោះ GitHub Webhook នឹង Trigger Jenkins ឱ្យទាញយកកូដ "
        "សាងសង់ Docker Image, បញ្ជូន Image ទៅកាន់ Amazon ECR រួច Deploy ដោយស្វ័យប្រវត្តិកាន់ Docker Swarm Cluster ដោយគ្មាន Downtime ឡើយ។"
    )

    # 2. Architecture
    add_section_h1("២. ស្ថាបត្យកម្មប្រព័ន្ធ (Architecture)")
    add_body_p(
        "ប្រព័ន្ធត្រូវបានរចនាឡើងជាលំហូរស្វ័យប្រវត្តិកម្មពីដើមដល់ចប់ (End-to-End Flow)៖\n"
        "• Developer Push កូដទៅ GitHub (https://github.com/Sievphor/cloud_home_work_2.git)\n"
        "• GitHub Webhook ផ្ញើសញ្ញា Trigger ទៅកាន់ Jenkins Server\n"
        "• Jenkins ធ្វើការ Checkout Source Code និង Login ECR តាមរយៈ aws ecr get-login-password\n"
        "• Jenkins Build Docker Image ជាមួយ Platform linux/amd64 និងដាក់ Tag តាម Build Number\n"
        "• Jenkins Push Image ទៅ Private Amazon ECR (908157891706.dkr.ecr.us-east-1.amazonaws.com/sievphor-homework2)\n"
        "• Jenkins ភ្ជាប់តាម SSH ទៅ Swarm Manager (3.91.30.20) និងដំណើរការ docker stack deploy\n"
        "• Docker Swarm Update Replicas ទាំង ៣ ចែកគ្នានៅលើ Manager និង Workers ដោយប្រើ Rolling Update។"
    )
    add_figure_box("រូបភាពទី ១៖ ស្ថាបត្យកម្ម CI/CD Pipeline លើ AWS EC2")

    # 3. Implementation Steps
    add_section_h1("៣. ការអនុវត្តការងារជាក់ស្តែង (Implementation Steps)")
    add_body_p(
        "ជំហានទី ១៖ បង្កើត GitHub Repository\n"
        "បានបង្កើត Repository ឈ្មោះ 'cloud_home_work_2' ដែលមានឯកសារ index.html, Dockerfile, Jenkinsfile, "
        "docker-stack.yml និង .gitignore ដើម្បីធានាសុវត្ថិភាពមិនឱ្យរុញ Credentials ចូល Git ឡើយ។"
    )
    add_figure_box("រូបភាពទី ២៖ GitHub Repository cloud_home_work_2")

    add_body_p(
        "ជំហានទី ២៖ សរសេរ Dockerfile និងតេស្តលើ Local\n"
        "បានប្រើប្រាស់ Base Image 'nginx:alpine' ដែលតូច និងមានសុវត្ថិភាព។ បានតេស្ត docker build និង run លើ local ឃើញថាកម្មវិធីឆ្លើយតបយ៉ាងត្រឹមត្រូវ។"
    )
    add_figure_box("រូបភាពទី ៣៖ Dockerfile និងដំណើរការ Build លើ Local")

    add_body_p(
        "ជំហានទី ៣៖ បង្កើត Amazon ECR Repository\n"
        "បានបង្កើត Private Repository ឈ្មោះ 'sievphor-homework2' ក្នុង Region us-east-1 "
        "(URI: 908157891706.dkr.ecr.us-east-1.amazonaws.com/sievphor-homework2)។"
    )
    add_figure_box("រូបភាពទី ៤៖ Amazon ECR Private Repository ផ្ទុក Docker Images")

    add_body_p(
        "ជំហានទី ៤៖ កំណត់រចនាសម្ព័ន្ធ Jenkins Pipeline & Webhook\n"
        "បានបង្កើត Jenkins Pipeline Job ដោយប្រើ Jenkinsfile (Pipeline as Code) និងបានកំណត់ GitHub Webhook "
        "ដើម្បីឱ្យ Pipeline ដំណើរការ Build ដោយស្វ័យប្រវត្តិតាមព្រឹត្តិការណ៍ Push Event។"
    )
    add_figure_box("រូបភាពទី ៥៖ Jenkins Pipeline Success Stages (Build, Push, Deploy)")

    add_body_p(
        "ជំហានទី ៥៖ បង្កើត Docker Swarm Cluster លើ AWS EC2\n"
        "បានបង្កើត EC2 Instances ចំនួន ៣ រួមមាន 1 Manager Node (3.91.30.20) និង 2 Worker Nodes (34.229.205.236, 54.160.200.116)។ "
        "តាមរយៈពាក្យបញ្ជា 'docker node ls' បង្ហាញថាក្នាំង Nodes ទាំង ៣ ស្ថិតក្នុងស្ថានភាព Ready និង Active។"
    )
    add_figure_box("រូបភាពទី ៦៖ Docker Swarm Cluster (docker node ls បង្ហាញ Manager & 2 Workers)")

    add_body_p(
        "ជំហានទី ៦៖ Deploy Application និងផ្ទៀងផ្ទាត់ Rolling Update\n"
        "Stack ត្រូវបានដាក់ពង្រាយជាមួយ Replicas ចំនួន ៣ ដោយប្រើ update_config (parallelism: 1, delay: 10s) "
        "ធានានូវ Zero Downtime។ ពិនិត្យឃើញ Service homework2_web ដំណើរការ 3/3 Replicas លើ Port 8898។"
    )
    add_figure_box("រូបភាពទី ៧៖ ស្ថានភាព Docker Service (docker service ls និង ps homework2_web)")

    add_body_p(
        "ជំហានទី ៧៖ ការចូលប្រើប្រាស់ផ្ទាល់លើ Browser (Live Verification)\n"
        "កម្មវិធីអាចចូលប្រើប្រាស់បានជាក់ស្តែងតាមរយៈ Public IP របស់គ្រប់ Nodes ទាំងអស់៖\n"
        "• Manager: http://3.91.30.20:8898\n"
        "• Worker 1: http://34.229.205.236:8898\n"
        "• Worker 2: http://54.160.200.116:8898"
    )
    add_figure_box("រូបភាពទី ៨៖ Web Application បង្ហាញនៅលើ Web Browser តាមរយៈ Port 8898")

    # 4. Challenges
    add_section_h1("៤. បញ្ហាប្រឈម និងដំណោះស្រាយ (Challenges & Solutions)")
    ch_table = doc.add_table(rows=4, cols=3)
    ch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["បញ្ហាប្រឈម", "មូលហេតុ", "ដំណោះស្រាយ"]
    for j, h in enumerate(headers):
        cell = ch_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.name = 'Khmer OS Battambang'
        set_cell_background(cell, "E2E8F0")

    challenges = [
        ("ERR_CONNECTION_TIMED_OUT លើ Browser", "Security Group មិនទាន់អនុញ្ញាត Port 8898", "បន្ថែម Inbound Rule អនុញ្ញាត Port 8898 ពី Anywhere (0.0.0.0/0)"),
        ("Worker មិនអាច Join Swarm បាន", "Port 2377 ត្រូវបាន Firewall បិទ", "បើក Port 2377, 7946, និង 4789 នៅក្នុង AWS EC2 Security Group"),
        ("Downtime អំឡុងពេល Deploy កូដថ្មី", "Container ទាំងអស់ត្រូវបិទព្រមគ្នា", "កំណត់ update_config: parallelism 1, delay 10s, order start-first")
    ]
    for i, row_data in enumerate(challenges):
        row = ch_table.rows[i + 1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.text = val
            cell.paragraphs[0].runs[0].font.name = 'Khmer OS Battambang'
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(cell, "F8FAFC" if i % 2 == 0 else "FFFFFF")

    # 5. Conclusion
    add_section_h1("៥. សេចក្តីសន្និដ្ឋាន (Conclusion)")
    add_body_p(
        "ការអនុវត្តកិច្ចការ Homework 2 នេះបានបង្ហាញពីសារៈសំខាន់នៃការកសាង CI/CD Pipeline លើ Cloud ដោយរួមបញ្ចូល GitHub Webhook, "
        "Jenkinsfile, Amazon ECR និង Docker Swarm Cluster។ ប្រព័ន្ធនេះជួយកាត់បន្ថយពេលវេលា ធានាសុវត្ថិភាព និងភាពអាចជឿជាក់បានកម្រិតខ្ពស់សម្រាប់ Production។"
    )

    output_path = os.path.join(os.path.dirname(__file__), "TONN-Sievphor-Documentation_homework2.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    create_report()
